from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from src.build_coach_index import build_index_rows
from src.corpus import SOURCE_METADATA_SCHEMA, IngestionDiagnostic, load_chunks
from src.generated_approval import read_review_workbook, write_review_workbook


def _write_pdf(path: Path, pages: list[str | None]) -> None:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
        )
        if text is not None:
            escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(str(path))


def _candidate(source_fields: dict[str, object], grounded_on: str) -> dict[str, object]:
    return {
        "id": "gen_0",
        "state": "NY",
        "section": "Practice Management",
        "difficulty": "medium",
        "question": "What should the architect do?",
        "options": ["A", "B", "C", "D"],
        "correctOption": "B",
        "explanation": "B follows the source.",
        "codeReference": "Synthetic source",
        "examWeight": 10,
        "topic": "Synthetic",
        "grounded_on": grounded_on,
        "repaired": False,
        "verdicts": {},
        "dup_sim": 0.1,
        "accepted": True,
        "reject_reasons": [],
        **source_fields,
    }


class CorpusProvenanceTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.corpus = Path(self.temp.name).resolve()

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_pdf_uses_physical_pages_and_empty_page_does_not_shift(self) -> None:
        repeated = "Repeated physical-page source text " * 4
        _write_pdf(self.corpus / "pages.pdf", [repeated, None, repeated])
        diagnostics: list[IngestionDiagnostic] = []

        chunks = load_chunks(
            min_len=1, max_len=1000, corpus_dir=self.corpus, diagnostics=diagnostics
        )

        self.assertEqual([chunk.source_page for chunk in chunks], [1, 3])
        self.assertEqual(chunks[0].text, chunks[1].text)
        self.assertNotEqual(chunks[0].source_chunk_id, chunks[1].source_chunk_id)
        self.assertEqual(len(diagnostics), 1)
        self.assertEqual(diagnostics[0].code, "no_extracted_text")
        self.assertEqual(diagnostics[0].source_page, 2)

    def test_same_basename_in_different_directories_has_distinct_identity(self) -> None:
        for directory in (self.corpus / "one", self.corpus / "two"):
            directory.mkdir()
            (directory / "same.txt").write_text("identical source text " * 8, encoding="utf-8")

        chunks = load_chunks(min_len=1, corpus_dir=self.corpus)

        self.assertEqual([chunk.source for chunk in chunks], ["same.txt", "same.txt"])
        self.assertEqual(
            [chunk.source_path for chunk in chunks], ["one/same.txt", "two/same.txt"]
        )
        self.assertNotEqual(chunks[0].source_document, chunks[1].source_document)
        self.assertNotEqual(chunks[0].source_chunk_id, chunks[1].source_chunk_id)

    def test_ids_are_stable_and_source_change_changes_identity(self) -> None:
        source = self.corpus / "stable.txt"
        source.write_text("stable source text " * 8, encoding="utf-8")
        first = load_chunks(min_len=1, corpus_dir=self.corpus)[0]
        repeated = load_chunks(min_len=1, corpus_dir=self.corpus)[0]
        self.assertEqual(first.source_document, repeated.source_document)
        self.assertEqual(first.source_chunk_id, repeated.source_chunk_id)
        self.assertEqual(first.source_sha256, repeated.source_sha256)

        source.write_text("changed source text " * 8, encoding="utf-8")
        changed = load_chunks(min_len=1, corpus_dir=self.corpus)[0]
        self.assertNotEqual(first.source_sha256, changed.source_sha256)
        self.assertNotEqual(first.source_document, changed.source_document)
        self.assertNotEqual(first.source_chunk_id, changed.source_chunk_id)

    def test_long_paragraph_respects_max_len(self) -> None:
        (self.corpus / "long.txt").write_text("word " * 180, encoding="utf-8")

        chunks = load_chunks(min_len=1, max_len=100, corpus_dir=self.corpus)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(1 <= len(chunk.text) <= 100 for chunk in chunks))

    def test_explicit_edition_only_and_non_pdf_page_is_not_applicable(self) -> None:
        notes = self.corpus / "notes"
        notes.mkdir()
        (notes / "guide_2099.md").write_text("# Heading\n" + "known text " * 12, encoding="utf-8")
        (self.corpus / "edition_2050.txt").write_text("unknown edition text " * 8, encoding="utf-8")
        (self.corpus / "source_metadata.json").write_text(
            json.dumps(
                {
                    "schema": SOURCE_METADATA_SCHEMA,
                    "sources": {
                        "notes/guide_2099.md": {"edition": "Explicit 2024", "revision": "R2"}
                    },
                }
            ),
            encoding="utf-8",
        )

        chunks = load_chunks(min_len=1, corpus_dir=self.corpus)
        by_path = {chunk.source_path: chunk for chunk in chunks}
        explicit = by_path["notes/guide_2099.md"]
        unknown = by_path["edition_2050.txt"]
        self.assertEqual(explicit.source_edition, "Explicit 2024")
        self.assertEqual(explicit.source_revision, "R2")
        self.assertIsNone(unknown.source_edition)
        self.assertIsNone(unknown.source_revision)
        self.assertIsNone(explicit.source_page)
        self.assertIn("source_page", explicit.source_not_applicable_metadata)
        self.assertIn("source_edition", unknown.source_missing_metadata)
        self.assertNotIn("2099", str(explicit.source_edition))
        self.assertNotIn("2050", str(unknown.source_edition))

    def test_empty_text_file_is_reported_without_ocr_or_chunk(self) -> None:
        (self.corpus / "empty.txt").write_bytes(b"")
        diagnostics: list[IngestionDiagnostic] = []

        chunks = load_chunks(min_len=1, corpus_dir=self.corpus, diagnostics=diagnostics)

        self.assertEqual(chunks, [])
        self.assertEqual([item.code for item in diagnostics], ["no_extracted_text"])
        self.assertIsNone(diagnostics[0].source_page)

    def test_metadata_survives_candidate_review_and_index_serialization(self) -> None:
        source = self.corpus / "docs" / "source.pdf"
        source.parent.mkdir()
        _write_pdf(source, ["traceable source text " * 10])
        chunk = load_chunks(min_len=1, corpus_dir=self.corpus)[0]
        source_fields = chunk.candidate_source_metadata()
        candidate = _candidate(source_fields, chunk.grounding_label())

        candidate_path = self.corpus / "candidate.json"
        candidate_path.write_text(json.dumps([candidate]), encoding="utf-8")
        workbook_path = self.corpus / "review.xlsx"
        write_review_workbook([candidate], workbook_path, source_path=candidate_path)
        workbook = load_workbook(workbook_path)
        sheet = workbook["review"]
        headings = {cell.value: cell.column for cell in sheet[1]}
        self.assertEqual(sheet.cell(2, headings["source_page"]).value, 1)
        self.assertEqual(
            sheet.cell(2, headings["source_chunk_id"]).value, chunk.source_chunk_id
        )
        sheet.cell(2, headings["REVIEW: verdict"], "Approve")
        workbook.save(workbook_path)
        workbook.close()
        review = read_review_workbook([candidate], workbook_path, source_path=candidate_path)
        self.assertTrue(review.decisions[0].approved)

        workbook = load_workbook(workbook_path)
        sheet = workbook["review"]
        headings = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(2, headings["source_page"], 2)
        workbook.save(workbook_path)
        workbook.close()
        changed_review = read_review_workbook(
            [candidate], workbook_path, source_path=candidate_path
        )
        self.assertFalse(changed_review.decisions[0].approved)
        self.assertIn("source_page", changed_review.decisions[0].reason or "")

        rows = build_index_rows([chunk])
        serialized = json.loads(json.dumps(rows))
        for key, value in source_fields.items():
            self.assertEqual(serialized[0][key], value)


if __name__ == "__main__":
    unittest.main()
