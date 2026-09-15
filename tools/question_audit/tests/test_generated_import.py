from __future__ import annotations

import copy
import json
import os
import tempfile
import threading
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.generated_approval import (
    ApprovalError,
    load_candidates,
    read_review_workbook,
    validate_candidates,
    write_review_workbook,
)
from src.generated_import import (
    ImportSafetyError,
    apply_import_plan,
    build_import_plan,
    load_bank,
    load_journal,
    pending_transaction_state,
    recover_pending_transaction,
    transaction_pointer_for,
    verify_bank_journal_consistency,
)
from src.merge_accepted import main as merge_main
from src.review_generated import main as review_main


def _candidate(candidate_id: str = "gen_0", question: str = "What should the architect do?"):
    return {
        "id": candidate_id,
        "state": "NY",
        "section": "Practice Management",
        "difficulty": "medium",
        "question": question,
        "options": ["A", "B", "C", "D"],
        "correctOption": "B",
        "explanation": "B follows the cited source.",
        "codeReference": "Example source, section 1",
        "examWeight": 10,
        "topic": "Contracts",
        "grounded_on": "example.pdf :: chunk_7",
        "repaired": False,
        "verdicts": {"consistency": {"pass": True}},
        "dup_sim": 0.12,
        "accepted": True,
        "reject_reasons": [],
        "generator_run_id": "synthetic-run-1",
    }


def _bank_row():
    return {
        "id": "pcm_1",
        "state": "NY",
        "section": "Practice Management",
        "difficulty": "easy",
        "question": "Existing question?",
        "options": ["One", "Two", "Three", "Four"],
        "correctOption": "One",
        "explanation": "Existing explanation.",
        "codeReference": "Existing source",
        "examWeight": 5,
        "topic": "Existing",
    }


class GeneratedImportTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        # Hosted Windows TEMP can use an 8.3 alias (RUNNER~1). Fault-injection
        # targets must identify the same resolved paths as the real importer.
        self.root = Path(self.temp.name).resolve()
        self.source = self.root / "generated_accepted.json"
        self.review = self.root / "generated_review.xlsx"
        self.bank = self.root / "questions_ny.json"
        self.journal = self.root / "questions_ny.provenance.json"
        self.backups = self.root / "backups"
        self.candidates = [_candidate()]
        self.source.write_text(json.dumps(self.candidates), encoding="utf-8")
        self.bank.write_text(json.dumps([_bank_row()]), encoding="utf-8")
        write_review_workbook(self.candidates, self.review, source_path=self.source)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def _edit_review(self, **values) -> None:
        workbook = load_workbook(self.review)
        sheet = workbook["review"]
        headings = {cell.value: cell.column for cell in sheet[1]}
        for heading, value in values.items():
            sheet.cell(row=2, column=headings[heading], value=value)
        workbook.save(self.review)
        workbook.close()

    def _plan(self):
        candidates = load_candidates(self.source)
        review = read_review_workbook(candidates, self.review, source_path=self.source)
        return build_import_plan(
            load_bank(self.bank),
            load_journal(self.journal),
            review,
            review_path=self.review,
            imported_at="2026-09-14T12:00:00+00:00",
        )

    def test_no_approval_and_omitted_reviewed_flag_cannot_add(self):
        before_bank = self.bank.read_bytes()
        before_review = self.review.read_bytes()
        result = merge_main(
            [
                "--input",
                str(self.source),
                "--review-book",
                str(self.review),
                "--bank",
                str(self.bank),
                "--journal",
                str(self.journal),
                "--backups-dir",
                str(self.backups),
                "--apply",
            ]
        )
        self.assertEqual(result, 0)
        self.assertEqual(self.bank.read_bytes(), before_bank)
        self.assertEqual(self.review.read_bytes(), before_review)
        self.assertFalse(self.journal.exists())
        self.assertFalse(self.backups.exists())
        self.assertEqual(len(self._plan().blocked), 1)

    def test_missing_and_legacy_review_workbooks_fail_closed(self):
        self.review.unlink()
        result = merge_main(
            [
                "--input",
                str(self.source),
                "--review-book",
                str(self.review),
                "--bank",
                str(self.bank),
                "--journal",
                str(self.journal),
            ]
        )
        self.assertEqual(result, 2)
        legacy = Workbook()
        legacy.active.append(["REVIEW: verdict", "id"])
        legacy.save(self.review)
        with self.assertRaisesRegex(ApprovalError, "legacy/invalid"):
            read_review_workbook(self.candidates, self.review, source_path=self.source)

    def test_reject_needs_edit_and_invalid_verdicts_are_blocked(self):
        for verdict in ("Reject", "Needs edit", "approved"):
            with self.subTest(verdict=verdict):
                self._edit_review(**{"REVIEW: verdict": verdict})
                plan = self._plan()
                self.assertEqual(len(plan.additions), 0)
                self.assertEqual(len(plan.blocked), 1)

    def test_approved_dry_run_changes_no_files(self):
        self._edit_review(**{"REVIEW: verdict": "Approve"})
        before_bank = self.bank.read_bytes()
        before_review = self.review.read_bytes()
        result = merge_main(
            [
                "--input",
                str(self.source),
                "--review-book",
                str(self.review),
                "--bank",
                str(self.bank),
                "--journal",
                str(self.journal),
                "--backups-dir",
                str(self.backups),
            ]
        )
        self.assertEqual(result, 0)
        self.assertEqual(self.bank.read_bytes(), before_bank)
        self.assertEqual(self.review.read_bytes(), before_review)
        self.assertFalse(self.journal.exists())
        self.assertFalse(self.backups.exists())

    def test_valid_approval_preserves_mapping_provenance_and_is_idempotent(self):
        self._edit_review(**{"REVIEW: verdict": "Approve", "REVIEW: notes": "Checked"})
        plan = self._plan()
        self.assertEqual(len(plan.additions), 1)
        transaction_dir = apply_import_plan(
            plan,
            bank_path=self.bank,
            journal_path=self.journal,
            backups_dir=self.backups,
        )
        self.assertIsNotNone(transaction_dir)
        bank = load_bank(self.bank)
        journal = load_journal(self.journal)
        self.assertEqual(len(bank), 2)
        self.assertEqual(bank[-1]["id"], "gen_q1")
        entry = journal["imports"][0]
        self.assertEqual(entry["candidate_id"], "gen_0")
        self.assertEqual(entry["bank_id"], "gen_q1")
        self.assertEqual(entry["candidate_snapshot"], self.candidates[0])
        self.assertEqual(entry["source_provenance"]["grounded_on"], "example.pdf :: chunk_7")
        self.assertFalse(entry["source_provenance"]["complete"])
        self.assertEqual(
            entry["source_provenance"]["missing_fields"],
            ["source_document", "source_edition", "source_page", "source_sha256"],
        )
        verify_bank_journal_consistency(bank, journal)

        repeated = self._plan()
        self.assertEqual(len(repeated.additions), 0)
        self.assertEqual(len(repeated.already_imported), 1)
        self.assertIsNone(
            apply_import_plan(
                repeated,
                bank_path=self.bank,
                journal_path=self.journal,
                backups_dir=self.backups,
            )
        )
        self.assertEqual(len(load_bank(self.bank)), 2)
        self.assertEqual(len(load_journal(self.journal)["imports"]), 1)

    def test_candidate_change_after_review_requires_new_approval(self):
        self._edit_review(**{"REVIEW: verdict": "Approve"})
        changed = copy.deepcopy(self.candidates)
        changed[0]["explanation"] = "Changed after review."
        self.source.write_text(json.dumps(changed), encoding="utf-8")
        with self.assertRaisesRegex(ApprovalError, "snapshot differs|source file hash differs"):
            read_review_workbook(changed, self.review, source_path=self.source)

    def test_visible_excel_change_is_blocked_even_with_hidden_fingerprint(self):
        self._edit_review(
            **{
                "REVIEW: verdict": "Approve",
                "question": "Manually changed only in Excel",
            }
        )
        plan = self._plan()
        self.assertEqual(len(plan.additions), 0)
        self.assertIn("review row differs in: question", plan.blocked[0].reason)

    def test_duplicate_candidate_and_review_ids_fail_closed(self):
        with self.assertRaisesRegex(ApprovalError, "duplicate/conflicting candidate id"):
            validate_candidates([self.candidates[0], copy.deepcopy(self.candidates[0])])

        workbook = load_workbook(self.review)
        sheet = workbook["review"]
        sheet.append([cell.value for cell in sheet[2]])
        workbook.save(self.review)
        workbook.close()
        with self.assertRaisesRegex(ApprovalError, "duplicate/conflicting review row id"):
            read_review_workbook(self.candidates, self.review, source_path=self.source)

    def test_conflicting_journal_ids_fail_closed(self):
        self._edit_review(**{"REVIEW: verdict": "Approve"})
        plan = self._plan()
        journal = plan.journal_after
        journal["imports"].append(copy.deepcopy(journal["imports"][0]))
        with self.assertRaisesRegex(ImportSafetyError, "duplicate/conflicting journal candidate id"):
            verify_bank_journal_consistency(plan.bank_after, journal)

    def test_review_generator_refuses_existing_output(self):
        before = self.review.read_bytes()
        result = review_main(["--input", str(self.source), "--output", str(self.review)])
        self.assertEqual(result, 2)
        self.assertEqual(self.review.read_bytes(), before)

    def test_bank_replace_failure_is_detected_and_recovered(self):
        self._assert_write_failure_is_recoverable(fail_target=self.bank)

    def test_journal_replace_failure_is_detected_and_recovered(self):
        self._assert_write_failure_is_recoverable(fail_target=self.journal)

    def test_recovery_refuses_target_outside_expected_before_after_hashes(self):
        _, pointer = self._leave_partial_transaction(fail_target=self.journal)
        unexpected = load_bank(self.bank) + [{**_bank_row(), "id": "external_change"}]
        self.bank.write_text(json.dumps(unexpected), encoding="utf-8")
        unexpected_bytes = self.bank.read_bytes()

        with self.assertRaisesRegex(ImportSafetyError, "neither before nor after hash"):
            recover_pending_transaction(
                pointer,
                expected_bank_path=self.bank,
                expected_journal_path=self.journal,
            )
        self.assertEqual(self.bank.read_bytes(), unexpected_bytes)
        self.assertFalse(self.journal.exists())
        self.assertTrue(pointer.exists())

    def test_dry_run_reports_partial_transaction_without_recovery_or_write(self):
        _, pointer = self._leave_partial_transaction(fail_target=self.journal)
        bank_after_failure = self.bank.read_bytes()
        pointer_before = pointer.read_bytes()

        result = merge_main(
            [
                "--input",
                str(self.source),
                "--review-book",
                str(self.review),
                "--bank",
                str(self.bank),
                "--journal",
                str(self.journal),
                "--backups-dir",
                str(self.backups),
            ]
        )
        self.assertEqual(result, 2)
        self.assertEqual(self.bank.read_bytes(), bank_after_failure)
        self.assertEqual(pointer.read_bytes(), pointer_before)
        self.assertFalse(self.journal.exists())

    def test_concurrent_apply_is_refused_while_transaction_is_pending(self):
        self._edit_review(**{"REVIEW: verdict": "Approve"})
        plan = self._plan()
        bank_replace_entered = threading.Event()
        release_bank_replace = threading.Event()
        thread_errors: list[BaseException] = []

        def paused_replace(source, destination):
            if Path(destination) == self.bank:
                bank_replace_entered.set()
                if not release_bank_replace.wait(timeout=10):
                    raise TimeoutError("test did not release bank replacement")
            return os.replace(source, destination)

        def first_apply():
            try:
                apply_import_plan(
                    plan,
                    bank_path=self.bank,
                    journal_path=self.journal,
                    backups_dir=self.backups,
                    replace_target=paused_replace,
                )
            except BaseException as exc:  # surfaced in the main test thread below
                thread_errors.append(exc)

        worker = threading.Thread(target=first_apply, daemon=True)
        worker.start()
        self.assertTrue(bank_replace_entered.wait(timeout=10))
        with self.assertRaisesRegex(ImportSafetyError, "bank import lock"):
            apply_import_plan(
                plan,
                bank_path=self.bank,
                journal_path=self.journal,
                backups_dir=self.backups,
            )
        release_bank_replace.set()
        worker.join(timeout=10)
        self.assertFalse(worker.is_alive())
        self.assertEqual(thread_errors, [])
        self.assertEqual(len(load_bank(self.bank)), 2)
        self.assertEqual(len(load_journal(self.journal)["imports"]), 1)

    def _assert_write_failure_is_recoverable(self, *, fail_target: Path) -> None:
        _, pointer = self._leave_partial_transaction(fail_target=fail_target)
        self.assertIsNotNone(pending_transaction_state(pointer))
        self.assertTrue(
            recover_pending_transaction(
                pointer,
                expected_bank_path=self.bank,
                expected_journal_path=self.journal,
            )
        )
        self.assertFalse(pointer.exists())
        bank = load_bank(self.bank)
        journal = load_journal(self.journal)
        verify_bank_journal_consistency(bank, journal)
        self.assertEqual(len(bank), 2)
        self.assertEqual(len(journal["imports"]), 1)
        entry = journal["imports"][0]
        self.assertEqual(entry["candidate_snapshot"], self.candidates[0])
        self.assertEqual(bank[-1]["question"], self.candidates[0]["question"])
        self.assertEqual(bank[-1]["options"], self.candidates[0]["options"])

    def _leave_partial_transaction(self, *, fail_target: Path):
        self._edit_review(**{"REVIEW: verdict": "Approve"})
        plan = self._plan()
        pointer = transaction_pointer_for(self.journal)
        failed = False

        def replace_once(source, destination):
            nonlocal failed
            if Path(destination) == fail_target and not failed:
                failed = True
                raise OSError("synthetic target write failure")
            return os.replace(source, destination)

        with self.assertRaisesRegex(OSError, "synthetic target write failure"):
            apply_import_plan(
                plan,
                bank_path=self.bank,
                journal_path=self.journal,
                backups_dir=self.backups,
                transaction_pointer=pointer,
                replace_target=replace_once,
            )
        return plan, pointer


if __name__ == "__main__":
    unittest.main()
