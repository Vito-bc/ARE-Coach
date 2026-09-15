"""Versioned human-review workbooks for generated question candidates.

The workbook is a review surface, not the source of truth.  Every row carries a
fingerprint of the complete candidate JSON and the visible cells are checked
again at import time.  Editing either the candidate or the reviewed row therefore
requires a newly generated workbook and a new human verdict.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from src.corpus import PAGE_AWARE_SOURCE_METADATA_FIELDS, SOURCE_METADATA_FIELDS


FINGERPRINT_SCHEMA = "are-coach.generated-candidate.v1"
WORKBOOK_SCHEMA = "are-coach.generated-review.v1"
PROVENANCE_WORKBOOK_SCHEMA = "are-coach.generated-review.v2"
POLICY_WORKBOOK_SCHEMA = "are-coach.generated-review.v3"
REVIEW_SHEET = "review"
METADATA_SHEET = "_metadata"
VERDICTS = ("Approve", "Reject", "Needs edit")

FINGERPRINT_COLUMN = "candidate_fingerprint"
REVIEW_COLUMNS = [
    "REVIEW: verdict",
    "REVIEW: notes",
    "id",
    FINGERPRINT_COLUMN,
    "section",
    "question",
    "options (>> = correct)",
    "correct",
    "explanation",
    "codeReference",
    "grounded_on",
    "repaired",
    "dup_sim",
]
VISIBLE_CANDIDATE_COLUMNS = [
    "id",
    "section",
    "question",
    "options (>> = correct)",
    "correct",
    "explanation",
    "codeReference",
    "grounded_on",
    "repaired",
    "dup_sim",
]
PROVENANCE_REVIEW_COLUMNS = list(PAGE_AWARE_SOURCE_METADATA_FIELDS)
POLICY_REVIEW_COLUMNS = [*SOURCE_METADATA_FIELDS, "source_policy_decision"]


class ApprovalError(ValueError):
    """The candidate snapshot or review workbook cannot be trusted."""


@dataclass(frozen=True)
class ReviewDecision:
    candidate: dict[str, Any]
    fingerprint: str
    verdict: str
    notes: str
    row_number: int
    approved: bool
    reason: str | None


@dataclass(frozen=True)
class ReviewResult:
    decisions: list[ReviewDecision]
    workbook_sha256: str
    source_sha256: str


def _canonical_json(value: Any) -> bytes:
    try:
        text = json.dumps(
            value,
            ensure_ascii=False,
            allow_nan=False,
            sort_keys=True,
            separators=(",", ":"),
        )
    except (TypeError, ValueError) as exc:
        raise ApprovalError(f"candidate is not canonical-JSON serializable: {exc}") from exc
    return text.encode("utf-8")


def sha256_bytes(data: bytes) -> str:
    return f"sha256:{hashlib.sha256(data).hexdigest()}"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return f"sha256:{digest.hexdigest()}"


def candidate_fingerprint(candidate: dict[str, Any]) -> str:
    """Hash every candidate field, including source/provenance metadata."""
    return sha256_bytes(
        _canonical_json({"schema": FINGERPRINT_SCHEMA, "candidate": candidate})
    )


def candidate_snapshot_fingerprint(candidates: list[dict[str, Any]]) -> str:
    return sha256_bytes(
        _canonical_json(
            {
                "schema": FINGERPRINT_SCHEMA,
                "candidate_fingerprints": [candidate_fingerprint(row) for row in candidates],
            }
        )
    )


def load_candidates(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise ApprovalError(f"candidate file does not exist: {path}")
    try:
        raw = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ApprovalError(f"cannot read candidate JSON {path}: {exc}") from exc
    return validate_candidates(raw)


def validate_candidates(raw: Any) -> list[dict[str, Any]]:
    if not isinstance(raw, list):
        raise ApprovalError("candidate JSON must be a top-level list")
    candidates: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, item in enumerate(raw):
        if not isinstance(item, dict):
            raise ApprovalError(f"candidate at index {index} is not an object")
        candidate_id = item.get("id")
        if not isinstance(candidate_id, str) or not candidate_id.strip():
            raise ApprovalError(f"candidate at index {index} has no non-empty string id")
        if candidate_id in seen:
            raise ApprovalError(f"duplicate/conflicting candidate id: {candidate_id}")
        # Fail before producing a fingerprint if a value cannot be reproduced.
        candidate_fingerprint(item)
        seen.add(candidate_id)
        candidates.append(item)
    return candidates


def review_row(candidate: dict[str, Any]) -> dict[str, Any]:
    options = candidate.get("options")
    correct = candidate.get("correctOption")
    if not isinstance(options, list):
        options_text = ""
    else:
        options_text = "\n".join(
            f"{'>>' if option == correct else '  '} {option}" for option in options
        )
    row = {
        "REVIEW: verdict": "",
        "REVIEW: notes": "",
        "id": candidate.get("id", ""),
        FINGERPRINT_COLUMN: candidate_fingerprint(candidate),
        "section": candidate.get("section", ""),
        "question": candidate.get("question", ""),
        "options (>> = correct)": options_text,
        "correct": correct if correct is not None else "",
        "explanation": candidate.get("explanation", ""),
        "codeReference": candidate.get("codeReference", ""),
        "grounded_on": candidate.get("grounded_on", ""),
        "repaired": candidate.get("repaired", False),
        "dup_sim": candidate.get("dup_sim", ""),
    }
    for field in POLICY_REVIEW_COLUMNS:
        value = candidate.get(field, "")
        if isinstance(value, (dict, list, tuple)):
            value = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        row[field] = "" if value is None else value
    return row


def _workbook_contract(candidates: list[dict[str, Any]]) -> tuple[str, list[str], list[str]]:
    carries_policy = any(
        candidate.get("source_policy_schema") is not None
        or "source_policy_decision" in candidate
        for candidate in candidates
    )
    if carries_policy:
        return (
            POLICY_WORKBOOK_SCHEMA,
            [*REVIEW_COLUMNS, *POLICY_REVIEW_COLUMNS],
            [*VISIBLE_CANDIDATE_COLUMNS, *POLICY_REVIEW_COLUMNS],
        )
    carries_provenance = any(
        candidate.get("source_metadata_schema") is not None
        for candidate in candidates
    )
    if carries_provenance:
        return (
            PROVENANCE_WORKBOOK_SCHEMA,
            [*REVIEW_COLUMNS, *PROVENANCE_REVIEW_COLUMNS],
            [*VISIBLE_CANDIDATE_COLUMNS, *PROVENANCE_REVIEW_COLUMNS],
        )
    return WORKBOOK_SCHEMA, REVIEW_COLUMNS, VISIBLE_CANDIDATE_COLUMNS


def write_review_workbook(
    candidates: list[dict[str, Any]],
    output_path: Path,
    *,
    source_path: Path,
) -> None:
    """Create a new review workbook; existing files are never overwritten."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    candidates = validate_candidates(candidates)
    if not candidates:
        raise ApprovalError("no generated candidates to review")
    if output_path.exists():
        raise FileExistsError(
            f"review workbook already exists: {output_path}; choose a new --output path"
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook_schema, review_columns, _ = _workbook_contract(candidates)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = REVIEW_SHEET
    sheet.append(review_columns)
    for candidate in candidates:
        values = review_row(candidate)
        sheet.append([values[column] for column in review_columns])

    dark = PatternFill("solid", fgColor="1F2937")
    yellow = PatternFill("solid", fgColor="F2B807")
    wrap_top = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.fill = dark
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column_number in (1, 2):
        sheet.cell(row=1, column=column_number).fill = yellow
        sheet.cell(row=1, column=column_number).font = Font(bold=True, color="1F2937")

    widths = {
        "REVIEW: verdict": 16,
        "REVIEW: notes": 34,
        "id": 16,
        FINGERPRINT_COLUMN: 12,
        "section": 20,
        "question": 60,
        "options (>> = correct)": 52,
        "correct": 38,
        "explanation": 58,
        "codeReference": 34,
        "grounded_on": 30,
        "repaired": 9,
        "dup_sim": 8,
        "source_metadata_schema": 24,
        "source_document": 24,
        "source_sha256": 24,
        "source_path": 34,
        "source_page": 12,
        "source_locator": 28,
        "source_chunk_id": 24,
        "source_extraction_version": 34,
        "source_edition": 20,
        "source_revision": 20,
        "source_missing_metadata": 30,
        "source_not_applicable_metadata": 30,
        "source_policy_schema": 24,
        "source_family_id": 24,
        "source_title": 30,
        "source_issuing_authority": 24,
        "source_jurisdictions": 24,
        "source_scope": 30,
        "source_exam_divisions": 34,
        "source_applicability_status": 20,
        "source_applicability_evidence": 34,
        "source_usage_permission_status": 20,
        "source_permitted_uses": 34,
        "source_usage_permission_note": 34,
        "source_policy_profiles": 30,
        "source_policy_decision": 42,
    }
    wrap_columns = {
        "REVIEW: notes",
        "question",
        "options (>> = correct)",
        "correct",
        "explanation",
        "codeReference",
        "grounded_on",
        *POLICY_REVIEW_COLUMNS,
    }
    last_row = len(candidates) + 1
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(review_columns))}{last_row}"
    for number, heading in enumerate(review_columns, 1):
        letter = get_column_letter(number)
        sheet.column_dimensions[letter].width = widths.get(heading, 16)
        if heading in wrap_columns:
            for cell in sheet[letter][1:]:
                cell.alignment = wrap_top
    fingerprint_letter = get_column_letter(REVIEW_COLUMNS.index(FINGERPRINT_COLUMN) + 1)
    sheet.column_dimensions[fingerprint_letter].hidden = True

    helper_column = len(review_columns) + 3
    helper_letter = get_column_letter(helper_column)
    for index, verdict in enumerate(VERDICTS, start=2):
        sheet.cell(row=index, column=helper_column, value=verdict)
    sheet.column_dimensions[helper_letter].hidden = True
    validation = DataValidation(
        type="list",
        formula1=f"${helper_letter}$2:${helper_letter}${1 + len(VERDICTS)}",
        allow_blank=True,
    )
    validation.add(f"A2:A{last_row}")
    sheet.add_data_validation(validation)

    metadata = workbook.create_sheet(METADATA_SHEET)
    metadata.sheet_state = "hidden"
    source_sha = sha256_file(source_path)
    metadata_rows = [
        ("workbook_schema", workbook_schema),
        ("fingerprint_schema", FINGERPRINT_SCHEMA),
        ("candidate_snapshot_fingerprint", candidate_snapshot_fingerprint(candidates)),
        ("candidate_count", len(candidates)),
        ("source_filename", source_path.name),
        ("source_sha256", source_sha),
    ]
    for row in metadata_rows:
        metadata.append(row)

    # Serialize first, then use exclusive creation. This closes the check/write
    # race and guarantees that an existing expert workbook cannot be replaced.
    buffer = io.BytesIO()
    workbook.save(buffer)
    with output_path.open("xb") as stream:
        stream.write(buffer.getvalue())
        stream.flush()
        os.fsync(stream.fileno())


def _metadata(sheet: Any) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in sheet.iter_rows(min_row=1, max_col=2, values_only=True):
        if key is not None:
            result[str(key)] = value
    return result


def _cell_value(value: Any) -> Any:
    return "" if value is None else value


def read_review_workbook(
    candidates: list[dict[str, Any]],
    workbook_path: Path,
    *,
    source_path: Path,
) -> ReviewResult:
    """Validate workbook identity and return explicit per-row decisions."""
    from openpyxl import load_workbook

    candidates = validate_candidates(candidates)
    if not workbook_path.exists():
        raise ApprovalError(
            f"review workbook is required and does not exist: {workbook_path}"
        )
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    except Exception as exc:
        raise ApprovalError(f"cannot read review workbook {workbook_path}: {exc}") from exc
    try:
        if REVIEW_SHEET not in workbook.sheetnames or METADATA_SHEET not in workbook.sheetnames:
            raise ApprovalError(
                "legacy/invalid review workbook: versioned review and _metadata sheets are required; "
                "generate a new workbook and obtain fresh human verdicts"
            )
        sheet = workbook[REVIEW_SHEET]
        metadata = _metadata(workbook[METADATA_SHEET])
        workbook_schema = metadata.get("workbook_schema")
        expected_schema, review_columns, visible_columns = _workbook_contract(candidates)
        if workbook_schema not in (
            WORKBOOK_SCHEMA,
            PROVENANCE_WORKBOOK_SCHEMA,
            POLICY_WORKBOOK_SCHEMA,
        ):
            raise ApprovalError("unsupported or missing review workbook schema")
        if workbook_schema != expected_schema:
            raise ApprovalError(
                "review workbook schema does not match candidate provenance; create a new "
                "workbook and obtain fresh approval"
            )
        if metadata.get("fingerprint_schema") != FINGERPRINT_SCHEMA:
            raise ApprovalError("unsupported or missing candidate fingerprint schema")
        expected_snapshot = candidate_snapshot_fingerprint(candidates)
        if metadata.get("candidate_snapshot_fingerprint") != expected_snapshot:
            raise ApprovalError(
                "candidate snapshot differs from the one sent for review; create a new workbook "
                "and obtain fresh approval"
            )
        if metadata.get("candidate_count") != len(candidates):
            raise ApprovalError("review workbook candidate count does not match the candidate snapshot")
        expected_source_sha = sha256_file(source_path)
        if metadata.get("source_sha256") != expected_source_sha:
            raise ApprovalError("candidate source file hash differs from the reviewed snapshot")

        try:
            header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        except StopIteration as exc:
            raise ApprovalError("review sheet is empty") from exc
        headers = [_cell_value(value) for value in header_values]
        for required in review_columns:
            if headers.count(required) != 1:
                raise ApprovalError(f"review sheet must contain exactly one {required!r} column")
        indexes = {heading: headers.index(heading) for heading in review_columns}

        rows_by_id: dict[str, tuple[int, tuple[Any, ...]]] = {}
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            candidate_id = _cell_value(row[indexes["id"]] if indexes["id"] < len(row) else None)
            if candidate_id == "":
                continue
            candidate_id = str(candidate_id)
            if candidate_id in rows_by_id:
                raise ApprovalError(f"duplicate/conflicting review row id: {candidate_id}")
            rows_by_id[candidate_id] = (row_number, row)

        expected_ids = {str(candidate["id"]) for candidate in candidates}
        actual_ids = set(rows_by_id)
        if actual_ids != expected_ids:
            missing = sorted(expected_ids - actual_ids)
            extra = sorted(actual_ids - expected_ids)
            raise ApprovalError(f"review row IDs do not match snapshot; missing={missing}, extra={extra}")

        decisions: list[ReviewDecision] = []
        for candidate in candidates:
            candidate_id = str(candidate["id"])
            row_number, row = rows_by_id[candidate_id]

            def value(column: str) -> Any:
                index = indexes[column]
                return _cell_value(row[index] if index < len(row) else None)

            expected = review_row(candidate)
            mismatches = [
                column
                for column in visible_columns
                if value(column) != _cell_value(expected[column])
            ]
            expected_fingerprint = candidate_fingerprint(candidate)
            reasons: list[str] = []
            if value(FINGERPRINT_COLUMN) != expected_fingerprint:
                reasons.append("candidate fingerprint does not match")
            if mismatches:
                reasons.append("review row differs in: " + ", ".join(mismatches))

            verdict = str(value("REVIEW: verdict")).strip()
            notes = str(value("REVIEW: notes"))
            if verdict != "Approve":
                if verdict in ("", "Reject", "Needs edit"):
                    reasons.append(f"human verdict is {verdict or 'blank'}")
                else:
                    reasons.append(f"invalid human verdict {verdict!r}")
            decisions.append(
                ReviewDecision(
                    candidate=candidate,
                    fingerprint=expected_fingerprint,
                    verdict=verdict,
                    notes=notes,
                    row_number=row_number,
                    approved=not reasons,
                    reason="; ".join(reasons) if reasons else None,
                )
            )

        return ReviewResult(
            decisions=decisions,
            workbook_sha256=sha256_file(workbook_path),
            source_sha256=expected_source_sha,
        )
    finally:
        workbook.close()
